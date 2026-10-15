"""สร้างเอกสารทดสอบภาษาไทยหลายแบบ (ตารางมีเส้น/พื้นหลัง/ผสานช่อง, แบบฟอร์ม, รายการ, ตารางข้ามหน้า)
แล้วแปลงเป็น PDF ด้วย LibreOffice เพื่อใช้ทดสอบการแปลง PDF -> Word

รัน: python tests/make_samples.py <โฟลเดอร์ปลายทาง>
"""
import os
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TH = "TH Sarabun New"


def style_doc(doc, font=TH, size=16):
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), font)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(size * 2))
    rpr.append(szcs)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def bold_cs(run):
    run.bold = True
    run._r.get_or_add_rPr().append(OxmlElement("w:bCs"))


def borders(cell, **sides):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side, (val, sz) in sides.items():
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "000000")
        b.append(e)
    tcpr.append(b)


def sample_grid(path):
    doc = Document()
    style_doc(doc)
    p = doc.add_paragraph()
    bold_cs(p.add_run("ตารางสรุปค่าธรรมเนียมการขอรับใบอนุญาตขับรถ"))
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["ลำดับ", "รายการ", "ค่าธรรมเนียม (บาท)", "หมายเหตุ"]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = ""
        bold_cs(c.paragraphs[0].add_run(h))
        shade(c, "D9E2F3")
    rows = [("๑", "ใบอนุญาตขับรถยนต์ส่วนบุคคลชั่วคราว", "๑๐๐", "อายุ ๒ ปี"),
            ("๒", "ใบอนุญาตขับรถจักรยานยนต์ส่วนบุคคลชั่วคราว", "๕๐", "อายุ ๒ ปี"),
            ("๓", "ค่าคำขอ", "๕", "-"),
            ("๔", "ค่าทดสอบสมรรถภาพร่างกาย ตาบอดสี สายตาทางลึก ปฏิกิริยาเท้า", "๕๐", "จ่ายครั้งเดียว")]
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    # ผสานช่องแนวนอน
    cells = t.add_row().cells
    m = cells[0].merge(cells[2])
    m.text = "รวมทั้งสิ้น (ผสานช่อง)"
    cells[3].text = "๒๐๕"
    doc.add_paragraph()
    doc.add_paragraph("ตารางผสานแนวตั้ง")
    t2 = doc.add_table(rows=3, cols=3)
    t2.style = "Table Grid"
    t2.cell(0, 0).merge(t2.cell(2, 0)).text = "ผสานสามแถว"
    for r in range(3):
        t2.cell(r, 1).text = f"แถวที่ {r + 1}"
        t2.cell(r, 2).text = "ข้อมูล ทั้งนี้ ที่นี่ ปู่ น้ำ"
    doc.save(path)


def sample_form(path):
    doc = Document()
    style_doc(doc)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bold_cs(h.add_run("แบบคำขอรับใบอนุญาต"))
    t = doc.add_table(rows=4, cols=2)
    labels = ["ชื่อ-นามสกุล", "เลขประจำตัวประชาชน", "ที่อยู่", "โทรศัพท์"]
    for i, lb in enumerate(labels):
        t.cell(i, 0).text = lb + " :"
        t.cell(i, 1).text = " "
        borders(t.cell(i, 1), bottom=("dotted", 6))
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1)
    borders(box.cell(0, 0), top=("double", 6), bottom=("double", 6), left=("double", 6), right=("double", 6))
    box.cell(0, 0).text = "สำหรับเจ้าหน้าที่ : ตรวจสอบเอกสารครบถ้วนแล้ว ลงชื่อ ................................"
    doc.add_paragraph()
    doc.add_paragraph("ข้าพเจ้าขอรับรองว่าข้อความข้างต้นเป็นความจริงทุกประการ")
    doc.save(path)


def sample_lists(path):
    doc = Document()
    style_doc(doc)
    bold_cs(doc.add_paragraph().add_run("ขั้นตอนการดำเนินงาน"))
    for s in ["ยื่นคำขอพร้อมเอกสารหลักฐาน", "ตรวจสอบเอกสารและคุณสมบัติเบื้องต้น", "ทดสอบสมรรถภาพทางร่างกาย",
              "อบรมและทดสอบข้อเขียน (ผ่านเกณฑ์ร้อยละ ๙๐)", "ทดสอบภาคปฏิบัติ", "ชำระค่าธรรมเนียมและรับใบอนุญาต"]:
        doc.add_paragraph(s, style="List Number")
    doc.add_paragraph("หมายเหตุ :")
    for s in ["เอกสารต้องเป็นฉบับจริง", "รูปถ่ายขนาด ๑ นิ้ว จำนวน ๒ รูป", "Contact: www.dlt.go.th / 1584"]:
        doc.add_paragraph(s, style="List Bullet")
    p = doc.add_paragraph()
    p.add_run("ข้อความปกติ ")
    r = p.add_run("ตัวหนาสีแดง ")
    bold_cs(r)
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run("และ ")
    p.add_run("ตัวเอียง").italic = True
    doc.save(path)


def sample_long_table(path):
    doc = Document()
    style_doc(doc, size=14)
    doc.add_paragraph("รายชื่อผู้ผ่านการอบรม (ตารางยาวข้ามหน้า)")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["ลำดับ", "ชื่อ-นามสกุล", "หน่วยงาน", "ผลการทดสอบ"]):
        t.rows[0].cells[i].text = h
        shade(t.rows[0].cells[i], "E7E6E6")
    names = ["สมชาย ใจดี", "สมหญิง รักไทย", "ประเสริฐ มั่นคง", "วิไลวรรณ ศรีสุข", "ธนพงศ์ ชัยภูมิ"]
    for n in range(1, 61):
        c = t.add_row().cells
        c[0].text = str(n).translate(str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙"))
        c[1].text = names[n % 5]
        c[2].text = "สำนักงานขนส่งจังหวัด" + ["ชัยภูมิ", "ขอนแก่น", "นครราชสีมา"][n % 3]
        c[3].text = "ผ่าน" if n % 7 else "ไม่ผ่าน"
    doc.save(path)


def sample_sans(path):
    doc = Document()
    style_doc(doc, font="Tahoma", size=11)
    doc.add_paragraph("เอกสารฟอนต์ Tahoma (ไม่มีหัว) สำหรับทดสอบการเลือกฟอนต์")
    t = doc.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"ข้อมูล {r + 1}-{c + 1} ทั้งนี้"
    doc.save(path)


def sample_double(path):
    """ตารางกรอบนอกเส้นคู่ เส้นในเส้นเดี่ยว + หัวตารางมีพื้นหลัง (แบบฟอร์มราชการที่พบบ่อย)"""
    doc = Document()
    style_doc(doc)
    bold_cs(doc.add_paragraph().add_run("บัญชีรายการเอกสารประกอบคำขอ"))
    t = doc.add_table(rows=4, cols=3)
    for r in range(4):
        for c in range(3):
            cell = t.cell(r, c)
            cell.text = ["ลำดับ", "รายการเอกสาร", "จำนวน"][c] if r == 0 else \
                [str(r), ["สำเนาบัตรประชาชน", "ใบรับรองแพทย์", "รูปถ่าย ๑ นิ้ว"][r - 1], "๑ ฉบับ"][c]
            sides = {"top": ("double" if r == 0 else "single", 6), "bottom": ("double" if r == 3 else "single", 6),
                     "left": ("double" if c == 0 else "single", 6), "right": ("double" if c == 2 else "single", 6)}
            borders(cell, **sides)
            if r == 0:
                shade(cell, "D9D9D9")
    doc.add_paragraph()
    doc.add_paragraph("ลงชื่อ ........................................ ผู้ยื่นคำขอ")
    doc.save(path)


SAMPLES = {"s1_grid": sample_grid, "s2_form": sample_form, "s3_lists": sample_lists,
           "s4_long_table": sample_long_table, "s5_sans": sample_sans, "s6_double": sample_double}

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "samples"
    os.makedirs(out, exist_ok=True)
    for name, fn in SAMPLES.items():
        docx = os.path.join(out, name + ".docx")
        fn(docx)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out, docx],
                       capture_output=True, timeout=300)
        print("made", name)
